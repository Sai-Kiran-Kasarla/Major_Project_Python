import streamlit as st
import os
import hashlib
import random
import time
from datetime import datetime

import pandas as pd
import numpy as np
from openpyxl import Workbook, load_workbook


class LMS:

    FILE_NAME = "lms_data.xlsx"

    ADMIN_EMAIL = st.secrets["ADMIN_EMAIL"]
    ADMIN_PASSWORD = st.secrets["ADMIN_PASSWORD"]

    def __init__(self):

        self.create_database()
        self.ensure_module_link_column()
        self.ensure_course_completion_sheet()
        self.remove_duplicate_enrollments()

    # ------------------------------------------------
    # DATABASE
    # ------------------------------------------------

    def create_database(self):

        if not os.path.exists(self.FILE_NAME):

            workbook = Workbook()

            sheets = {
                "Users": [
                    "ID",
                    "Name",
                    "Email",
                    "PasswordHash",
                    "Role",
                    "Approved"
                ],

                "Courses": [
                    "CourseID",
                    "Title",
                    "Description",
                    "Fee",
                    "Days",
                    "InstructorID"
                ],

                "Enrollments": [
                    "EnrollmentID",
                    "StudentID",
                    "CourseID",
                    "Percentage",
                    "Discount",
                    "FinalFee",
                    "Status"
                ],

                "Modules": [
                    "ModuleID",
                    "CourseID",
                    "Title",
                    "Description",
                    "Link"
                ],

                "Tasks": [
                    "TaskID",
                    "ModuleID",
                    "Title",
                    "Description"
                ],

                "Assignments": [
                    "AssignmentID",
                    "CourseID",
                    "Title",
                    "Description",
                    "TotalMarks"
                ],

                "AssignmentScores": [
                    "AssignmentID",
                    "StudentID",
                    "Marks"
                ],

                "Exams": [
                    "ExamID",
                    "CourseID",
                    "Title",
                    "TotalMarks"
                ],

                "Questions": [
                    "QuestionID",
                    "ExamID",
                    "Question",
                    "OptionA",
                    "OptionB",
                    "OptionC",
                    "OptionD",
                    "Answer"
                ],

                "ExamResults": [
                    "ExamID",
                    "StudentID",
                    "Marks",
                    "TotalMarks",
                    "Date"
                ],

                "Attendance": [
                    "AttendanceID",
                    "CourseID",
                    "StudentID",
                    "Date",
                    "Status"
                ],

                "Progress": [
                    "StudentID",
                    "CourseID",
                    "ModuleID",
                    "Completed"
                ],

                "CourseCompletions": [
                    "CompletionID",
                    "StudentID",
                    "CourseID",
                    "CompletedBy",
                    "CompletedDate",
                    "CertificateID"
                ]
            }

            first_sheet = True

            for sheet_name, headers in sheets.items():

                if first_sheet:

                    sheet = workbook.active
                    sheet.title = sheet_name

                    first_sheet = False

                else:

                    sheet = workbook.create_sheet(
                        sheet_name
                    )

                sheet.append(headers)

            workbook.save(self.FILE_NAME)

    # ------------------------------------------------
    # MODULE LINK MIGRATION
    # ------------------------------------------------

    def ensure_module_link_column(self):

        try:
            workbook = load_workbook(self.FILE_NAME)

            if "Modules" not in workbook.sheetnames:
                workbook.save(self.FILE_NAME)
                return

            sheet = workbook["Modules"]

            headers = []
            for cell in sheet[1]:
                headers.append(str(cell.value).strip() if cell.value is not None else "")

            if "Link" not in headers:
                sheet.cell(1, len(headers) + 1).value = "Link"

                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row, len(headers) + 1).value = ""

                workbook.save(self.FILE_NAME)

        except Exception as error:
            print("Module link migration error:", error)

    # ------------------------------------------------
    # GENERAL EXCEL FUNCTIONS
    # ------------------------------------------------

    def read_sheet(self, sheet_name):

        try:

            data = pd.read_excel(
                self.FILE_NAME,
                sheet_name=sheet_name
            )

            aliases = {
                "Course ID": "CourseID",
                "Course Id": "CourseID",
                "course_id": "CourseID",
                "COURSEID": "CourseID",
                "Student ID": "StudentID",
                "Student Id": "StudentID",
                "Instructor ID": "InstructorID",
                "Instructor Id": "InstructorID",
                "Enrollment ID": "EnrollmentID",
                "Module ID": "ModuleID",
                "Task ID": "TaskID",
                "Assignment ID": "AssignmentID",
                "Exam ID": "ExamID",
                "Question ID": "QuestionID",
                "Attendance ID": "AttendanceID"
            }

            data = data.rename(columns=aliases)

            expected = {
                "Users": ["ID", "Name", "Email", "PasswordHash", "Role", "Approved"],
                "Courses": ["CourseID", "Title", "Description", "Fee", "Days", "InstructorID"],
                "Enrollments": ["EnrollmentID", "StudentID", "CourseID", "Percentage", "Discount", "FinalFee", "Status"],
                "Modules": ["ModuleID", "CourseID", "Title", "Description", "Link"],
                "Tasks": ["TaskID", "ModuleID", "Title", "Description"],
                "Assignments": ["AssignmentID", "CourseID", "Title", "Description", "TotalMarks"],
                "AssignmentScores": ["AssignmentID", "StudentID", "Marks"],
                "Exams": ["ExamID", "CourseID", "Title", "TotalMarks"],
                "Questions": ["QuestionID", "ExamID", "Question", "OptionA", "OptionB", "OptionC", "OptionD", "Answer"],
                "ExamResults": ["ExamID", "StudentID", "Marks", "TotalMarks", "Date"],
                "Attendance": ["AttendanceID", "CourseID", "StudentID", "Date", "Status"],
                "Progress": ["StudentID", "CourseID", "ModuleID", "Completed"],
                "CourseCompletions": ["CompletionID", "StudentID", "CourseID", "CompletedBy", "CompletedDate", "CertificateID"]
            }

            required = expected.get(sheet_name, [])
            if required and not all(column in data.columns for column in required):
                if len(data.columns) == len(required):
                    data.columns = required

            return data

        except Exception as error:

            print("Excel load error:", error)

            return pd.DataFrame()

    def append_row(self, sheet_name, values):

        workbook = load_workbook(
            self.FILE_NAME
        )

        sheet = workbook[sheet_name]

        sheet.append(values)

        workbook.save(
            self.FILE_NAME
        )

    # ------------------------------------------------
    # PASSWORD
    # ------------------------------------------------

    def hash_password(self, password):

        return hashlib.sha256(
            password.encode()
        ).hexdigest()

    def check_password(
        self,
        password,
        password_hash
    ):

        return (
            self.hash_password(password)
            == password_hash
        )

    # ------------------------------------------------
    # ID GENERATOR
    # ------------------------------------------------

    def generate_id(
        self,
        prefix,
        sheet_name,
        column
    ):

        data = self.read_sheet(
            sheet_name
        )

        if data.empty:
            return prefix + "001"

        if column not in data.columns:

            if column == "CourseID" and "Course ID" in data.columns:
                column = "Course ID"

            elif column == "StudentID" and "Student ID" in data.columns:
                column = "Student ID"

            elif column == "InstructorID" and "Instructor ID" in data.columns:
                column = "Instructor ID"

            else:
                return prefix + "001"

        highest = 0

        for value in data[column]:

            text = str(value).strip()

            if text == "" or text.lower() == "nan":
                continue

            number_text = ""

            for character in text[::-1]:

                if character.isdigit():
                    number_text = character + number_text
                else:
                    break

            if number_text == "":
                continue

            try:
                number = int(number_text)

                if number > highest:
                    highest = number

            except ValueError:
                pass

        return prefix + str(
            highest + 1
        ).zfill(3)

    # ------------------------------------------------
    # DISCOUNT
    # ------------------------------------------------

    def calculate_discount(
        self,
        percentage
    ):

        percentage = float(
            percentage
        )

        if percentage >= 90:
            return 30

        if percentage >= 80:
            return 20

        if percentage >= 70:
            return 10

        if percentage >= 60:
            return 5

        return 0

    # ------------------------------------------------
    # USERS
    # ------------------------------------------------

    def register_user(
        self,
        name,
        email,
        password,
        role
    ):

        users = self.read_sheet(
            "Users"
        )

        email = email.strip().lower()

        if not users.empty:

            for old_email in users["Email"]:

                if str(
                    old_email
                ).lower() == email:

                    return False, "Email already registered."

        user_id = self.generate_id(
            "USR",
            "Users",
            "ID"
        )

        password_hash = self.hash_password(
            password
        )

        approved = True

        if role == "Instructor":
            approved = False

        self.append_row(
            "Users",
            [
                user_id,
                name,
                email,
                password_hash,
                role,
                approved
            ]
        )

        return True, user_id

    def get_user(
        self,
        email,
        role=None
    ):

        users = self.read_sheet(
            "Users"
        )

        if users.empty:
            return None

        email = email.strip().lower()

        for _, row in users.iterrows():

            row_email = str(
                row["Email"]
            ).lower()

            row_role = str(
                row["Role"]
            )

            if row_email == email:

                if role is None or row_role == role:

                    return row

        return None

    def login_user(
        self,
        email,
        password,
        role
    ):

        if role == "Admin":

            if (
                email.lower()
                == self.ADMIN_EMAIL.lower()
                and password
                == self.ADMIN_PASSWORD
            ):

                return {
                    "ID": "A001",
                    "Name": "Administrator",
                    "Email": self.ADMIN_EMAIL,
                    "Role": "Admin",
                    "Approved": True
                }

            return None

        user = self.get_user(
            email,
            role
        )

        if user is None:
            return None

        if not self.check_password(
            password,
            str(user["PasswordHash"])
        ):

            return None

        return user

    # ------------------------------------------------
    # APPROVAL
    # ------------------------------------------------

    def get_pending_users(
        self,
        role
    ):

        users = self.read_sheet(
            "Users"
        )

        if users.empty:
            return users

        return users[
            (users["Role"] == role)
            & (users["Approved"] == False)
        ]

    def approve_user(
        self,
        user_id
    ):

        workbook = load_workbook(
            self.FILE_NAME
        )

        sheet = workbook["Users"]

        for row in range(
            2,
            sheet.max_row + 1
        ):

            if str(
                sheet.cell(row, 1).value
            ) == str(user_id):

                sheet.cell(
                    row,
                    6
                ).value = True

        workbook.save(
            self.FILE_NAME
        )

    # ------------------------------------------------
    # COURSES
    # ------------------------------------------------

    def create_course(
        self,
        title,
        description,
        fee,
        days,
        instructor_id=""
    ):

        course_id = self.generate_id(
            "CRS",
            "Courses",
            "CourseID"
        )

        self.append_row(
            "Courses",
            [
                course_id,
                title,
                description,
                float(fee),
                int(days),
                instructor_id
            ]
        )

        return course_id

    def get_courses(self):

        return self.read_sheet(
            "Courses"
        )

    def assign_instructor(
        self,
        course_id,
        instructor_id
    ):

        workbook = load_workbook(
            self.FILE_NAME
        )

        sheet = workbook["Courses"]

        for row in range(
            2,
            sheet.max_row + 1
        ):

            if str(
                sheet.cell(row, 1).value
            ) == str(course_id):

                sheet.cell(
                    row,
                    6
                ).value = instructor_id

        workbook.save(
            self.FILE_NAME
        )

    # ------------------------------------------------
    # ENROLLMENT
    # ------------------------------------------------

    # ------------------------------------------------
    # REMOVE DUPLICATE ENROLLMENTS
    # ------------------------------------------------

    def remove_duplicate_enrollments(self):

        try:
            workbook = load_workbook(self.FILE_NAME)
            sheet = workbook["Enrollments"]

            if sheet.max_row <= 2:
                return

            seen = {}
            rows_to_delete = []

            # If duplicates already exist, keep one record for each
            # StudentID + CourseID combination. Prefer Assigned over Pending.
            for row in range(2, sheet.max_row + 1):
                student_id = str(sheet.cell(row, 2).value or "").strip()
                course_id = str(sheet.cell(row, 3).value or "").strip()
                status = str(sheet.cell(row, 7).value or "").strip()

                if not student_id or not course_id:
                    continue

                key = student_id + "|" + course_id

                if key not in seen:
                    seen[key] = row
                else:
                    old_row = seen[key]
                    old_status = str(sheet.cell(old_row, 7).value or "").strip()

                    if status.lower() == "assigned" and old_status.lower() != "assigned":
                        rows_to_delete.append(old_row)
                        seen[key] = row
                    else:
                        rows_to_delete.append(row)

            for row in sorted(rows_to_delete, reverse=True):
                sheet.delete_rows(row)

            if rows_to_delete:
                workbook.save(self.FILE_NAME)

        except Exception as error:
            print("Duplicate enrollment cleanup warning:", error)

    def request_enrollment(
        self,
        student_id,
        course_id,
        percentage
    ):

        courses = self.get_courses()

        if courses.empty:
            return False, "Course not found."

        course = courses[
            courses["CourseID"].astype(str)
            == str(course_id)
        ]

        if course.empty:
            return False, "Course not found."

        # A student can enroll in many different courses, but the same
        # student cannot enroll in the same course more than once.
        existing = self.get_student_enrollments(student_id)

        if not existing.empty and "CourseID" in existing.columns:
            same_course = existing[
                existing["CourseID"].astype(str)
                == str(course_id)
            ]

            if not same_course.empty:
                status = str(same_course.iloc[0].get("Status", ""))

                if status.lower() == "assigned":
                    return False, "You are already enrolled in this course."

                return False, "You already have an enrollment request for this course."

        fee = float(
            course.iloc[0]["Fee"]
        )

        discount = self.calculate_discount(
            percentage
        )

        final_fee = fee - (
            fee * discount / 100
        )

        enrollment_id = self.generate_id(
            "ENR",
            "Enrollments",
            "EnrollmentID"
        )

        self.append_row(
            "Enrollments",
            [
                enrollment_id,
                student_id,
                course_id,
                float(percentage),
                discount,
                final_fee,
                "Pending"
            ]
        )

        return True, final_fee

    def get_student_enrollments(
        self,
        student_id
    ):

        data = self.read_sheet(
            "Enrollments"
        )

        if data.empty:
            return data

        return data[
            data["StudentID"].astype(str)
            == str(student_id)
        ]

    def get_course_enrollments(
        self,
        course_id
    ):

        data = self.read_sheet(
            "Enrollments"
        )

        if data.empty:
            return data

        return data[
            data["CourseID"].astype(str)
            == str(course_id)
        ]

    def update_enrollment_status(
        self,
        enrollment_id,
        status
    ):

        workbook = load_workbook(
            self.FILE_NAME
        )

        sheet = workbook["Enrollments"]

        for row in range(
            2,
            sheet.max_row + 1
        ):

            if str(
                sheet.cell(row, 1).value
            ) == str(enrollment_id):

                sheet.cell(
                    row,
                    7
                ).value = status

        workbook.save(
            self.FILE_NAME
        )

    # ------------------------------------------------
    # MODULE
    # ------------------------------------------------

    def add_module(
        self,
        course_id,
        title,
        description,
        link=""
    ):

        module_id = self.generate_id(
            "MOD",
            "Modules",
            "ModuleID"
        )

        self.append_row(
            "Modules",
            [
                module_id,
                course_id,
                title,
                description,
                link
            ]
        )

        return module_id

    def get_modules(
        self,
        course_id
    ):

        data = self.read_sheet(
            "Modules"
        )

        if data.empty:
            return data

        return data[
            data["CourseID"].astype(str)
            == str(course_id)
        ]

    # ------------------------------------------------
    # TASK
    # ------------------------------------------------

    def add_task(
        self,
        module_id,
        title,
        description
    ):

        task_id = self.generate_id(
            "TSK",
            "Tasks",
            "TaskID"
        )

        self.append_row(
            "Tasks",
            [
                task_id,
                module_id,
                title,
                description
            ]
        )

        return task_id

    # ------------------------------------------------
    # ASSIGNMENT
    # ------------------------------------------------

    def add_assignment(
        self,
        course_id,
        title,
        description,
        total_marks
    ):

        assignment_id = self.generate_id(
            "ASN",
            "Assignments",
            "AssignmentID"
        )

        self.append_row(
            "Assignments",
            [
                assignment_id,
                course_id,
                title,
                description,
                total_marks
            ]
        )

        return assignment_id

    def add_assignment_score(
        self,
        assignment_id,
        student_id,
        marks
    ):

        self.append_row(
            "AssignmentScores",
            [
                assignment_id,
                student_id,
                marks
            ]
        )

    def get_assignment_scores(
        self,
        student_id
    ):

        data = self.read_sheet(
            "AssignmentScores"
        )

        if data.empty:
            return data

        return data[
            data["StudentID"].astype(str)
            == str(student_id)
        ]

    # ------------------------------------------------
    # EXAMS
    # ------------------------------------------------

    def create_exam(
        self,
        course_id,
        title,
        total_marks
    ):

        exam_id = self.generate_id(
            "EXM",
            "Exams",
            "ExamID"
        )

        self.append_row(
            "Exams",
            [
                exam_id,
                course_id,
                title,
                total_marks
            ]
        )

        return exam_id

    def add_question(
        self,
        exam_id,
        question,
        option_a,
        option_b,
        option_c,
        option_d,
        answer
    ):
        # Do not add the same question twice to one exam.
        existing = self.get_questions(exam_id)

        if not existing.empty and "Question" in existing.columns:
            wanted = str(question).strip().lower()
            for old_question in existing["Question"]:
                if str(old_question).strip().lower() == wanted:
                    return False

        question_id = self.generate_id(
            "QUE",
            "Questions",
            "QuestionID"
        )

        self.append_row(
            "Questions",
            [
                question_id,
                exam_id,
                question,
                option_a,
                option_b,
                option_c,
                option_d,
                answer
            ]
        )

        return True

    def get_questions(
        self,
        exam_id
    ):

        data = self.read_sheet(
            "Questions"
        )

        if data.empty:
            return data

        return data[
            data["ExamID"].astype(str)
            == str(exam_id)
        ]

    def save_exam_result(
        self,
        exam_id,
        student_id,
        marks,
        total_marks
    ):

        self.append_row(
            "ExamResults",
            [
                exam_id,
                student_id,
                marks,
                total_marks,
                datetime.now().strftime(
                    "%Y-%m-%d"
                )
            ]
        )

    def get_exam_results(
        self,
        student_id=None
    ):

        data = self.read_sheet(
            "ExamResults"
        )

        if data.empty:
            return data

        if student_id is not None:

            return data[
                data["StudentID"].astype(str)
                == str(student_id)
            ]

        return data

    # ------------------------------------------------
    # ATTENDANCE
    # ------------------------------------------------

    def mark_attendance(
        self,
        course_id,
        student_id,
        date,
        status
    ):

        data = self.read_sheet(
            "Attendance"
        )

        if not data.empty:

            existing = data[
                (data["CourseID"].astype(str)
                 == str(course_id))
                &
                (data["StudentID"].astype(str)
                 == str(student_id))
                &
                (data["Date"].astype(str)
                 == str(date))
            ]

            if not existing.empty:

                workbook = load_workbook(
                    self.FILE_NAME
                )

                sheet = workbook["Attendance"]

                for row in range(
                    2,
                    sheet.max_row + 1
                ):

                    same_course = (
                        str(
                            sheet.cell(
                                row,
                                2
                            ).value
                        )
                        == str(course_id)
                    )

                    same_student = (
                        str(
                            sheet.cell(
                                row,
                                3
                            ).value
                        )
                        == str(student_id)
                    )

                    same_date = (
                        str(
                            sheet.cell(
                                row,
                                4
                            ).value
                        )
                        == str(date)
                    )

                    if (
                        same_course
                        and same_student
                        and same_date
                    ):

                        sheet.cell(
                            row,
                            5
                        ).value = status

                workbook.save(
                    self.FILE_NAME
                )

                return

        attendance_id = self.generate_id(
            "ATT",
            "Attendance",
            "AttendanceID"
        )

        self.append_row(
            "Attendance",
            [
                attendance_id,
                course_id,
                student_id,
                date,
                status
            ]
        )

    def get_attendance(
        self,
        student_id=None,
        course_id=None
    ):

        data = self.read_sheet(
            "Attendance"
        )

        if data.empty:
            return data

        if student_id is not None:

            data = data[
                data["StudentID"].astype(str)
                == str(student_id)
            ]

        if course_id is not None:

            data = data[
                data["CourseID"].astype(str)
                == str(course_id)
            ]

        return data

    # ------------------------------------------------
    # COURSE COMPLETION / CERTIFICATES
    # ------------------------------------------------

    def mark_course_completed(self, student_id, course_id, instructor_id):
        """Mark a student's course as completed by an instructor."""
        data = self.read_sheet("CourseCompletions")
        if not data.empty:
            existing = data[(data["StudentID"].astype(str) == str(student_id)) &
                            (data["CourseID"].astype(str) == str(course_id))]
            if not existing.empty:
                return str(existing.iloc[0]["CertificateID"])

        certificate_id = self.generate_id("CERT", "CourseCompletions", "CompletionID")
        self.append_row("CourseCompletions", [
            certificate_id, student_id, course_id, instructor_id,
            datetime.now().strftime("%Y-%m-%d"), certificate_id
        ])
        return certificate_id

    def get_course_completions(self, student_id=None, course_id=None):
        """Return course completion records, optionally filtered."""
        data = self.read_sheet("CourseCompletions")
        if data.empty:
            return data
        if student_id is not None:
            data = data[data["StudentID"].astype(str) == str(student_id)]
        if course_id is not None:
            data = data[data["CourseID"].astype(str) == str(course_id)]
        return data

    def ensure_course_completion_sheet(self):
        """Create the course-completion sheet for existing databases."""
        try:
            workbook = load_workbook(self.FILE_NAME)
            if "CourseCompletions" not in workbook.sheetnames:
                sheet = workbook.create_sheet("CourseCompletions")
                sheet.append(["CompletionID", "StudentID", "CourseID", "CompletedBy", "CompletedDate", "CertificateID"])
                workbook.save(self.FILE_NAME)
        except Exception as error:
            print("Unable to create CourseCompletions sheet:", error)

    # ------------------------------------------------
    # PROGRESS
    # ------------------------------------------------

    def save_progress(
        self,
        student_id,
        course_id,
        module_id,
        completed
    ):

        data = self.read_sheet(
            "Progress"
        )

        if not data.empty:

            existing = data[
                (data["StudentID"].astype(str)
                 == str(student_id))
                &
                (data["CourseID"].astype(str)
                 == str(course_id))
                &
                (data["ModuleID"].astype(str)
                 == str(module_id))
            ]

            if not existing.empty:

                workbook = load_workbook(
                    self.FILE_NAME
                )

                sheet = workbook["Progress"]

                for row in range(
                    2,
                    sheet.max_row + 1
                ):

                    if (
                        str(
                            sheet.cell(
                                row,
                                1
                            ).value
                        ) == str(student_id)
                        and
                        str(
                            sheet.cell(
                                row,
                                2
                            ).value
                        ) == str(course_id)
                        and
                        str(
                            sheet.cell(
                                row,
                                3
                            ).value
                        ) == str(module_id)
                    ):

                        sheet.cell(
                            row,
                            4
                        ).value = completed

                workbook.save(
                    self.FILE_NAME
                )

                return

        self.append_row(
            "Progress",
            [
                student_id,
                course_id,
                module_id,
                completed
            ]
        )

    def get_progress(
        self,
        student_id,
        course_id
    ):

        data = self.read_sheet(
            "Progress"
        )

        if data.empty:
            return data

        return data[
            (data["StudentID"].astype(str)
             == str(student_id))
            &
            (data["CourseID"].astype(str)
             == str(course_id))
        ]

    # ------------------------------------------------
    # STATISTICS
    # ------------------------------------------------

    def get_statistics(self):

        students = self.read_sheet(
            "Users"
        )

        courses = self.read_sheet(
            "Courses"
        )

        enrollments = self.read_sheet(
            "Enrollments"
        )

        attendance = self.read_sheet(
            "Attendance"
        )

        if students.empty:
            student_count = 0
        else:
            student_count = len(
                students[
                    students["Role"]
                    == "Student"
                ]
            )

        if courses.empty:
            course_count = 0
        else:
            course_count = len(courses)

        if enrollments.empty:
            enrollment_count = 0
        else:
            enrollment_count = len(
                enrollments
            )

        if attendance.empty:

            attendance_rate = 0

        else:

            total = len(
                attendance
            )

            present = len(
                attendance[
                    attendance["Status"]
                    == "Present"
                ]
            )

            if total == 0:
                attendance_rate = 0
            else:
                attendance_rate = (
                    present / total
                ) * 100

        return {
            "students": student_count,
            "courses": course_count,
            "enrollments": enrollment_count,
            "attendance": round(
                attendance_rate,
                2
            )
        }

    def attendance_chart_data(
        self
    ):

        data = self.read_sheet(
            "Attendance"
        )

        if data.empty:

            return pd.DataFrame(
                {
                    "Status": [
                        "Present",
                        "Absent"
                    ],
                    "Count": [
                        0,
                        0
                    ]
                }
            )

        counts = data[
            "Status"
        ].value_counts()

        present = int(
            counts.get(
                "Present",
                0
            )
        )

        absent = int(
            counts.get(
                "Absent",
                0
            )
        )

        return pd.DataFrame(
            {
                "Status": [
                    "Present",
                    "Absent"
                ],
                "Count": [
                    present,
                    absent
                ]
            }
        )

    # ------------------------------------------------
    # STUDENT LIST FOR INSTRUCTOR
    # ------------------------------------------------

    def get_approved_students(self):

        users = self.read_sheet(
            "Users"
        )

        if users.empty:
            return users

        return users[
            (users["Role"] == "Student")
            &
            (users["Approved"] == True)
        ]

    def get_instructors(self):

        users = self.read_sheet(
            "Users"
        )

        if users.empty:
            return users

        return users[
            (users["Role"] == "Instructor")
            &
            (users["Approved"] == True)
        ]